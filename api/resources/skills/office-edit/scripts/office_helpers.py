"""Helpers for changing existing Word, PowerPoint and Excel files with python-docx, python-pptx and openpyxl."""

import copy
import hashlib
import os

import docx
import openpyxl
import pptx
from lxml import etree
from pptx.dml.color import RGBColor as PptxRGB

OUTPUT_DIR = "/output"  # where the sandbox collects the program's files

__all__ = [
    "set_text",
    "slide_shapes",
    "shape_by_role",
    "delete_shape",
    "delete_paragraph",
    "replace_text",
    "set_font",
    "set_text_color",
    "set_background",
    "slide_texts",
    "duplicate_slide",
    "delete_slide",
    "move_slide",
    "save",
]

# The file each document was opened from, with a fingerprint of its content, so that save() can
# tell when nothing was changed (a model replacing text that is not there, or changing nothing).
# The sandbox imports this module before the program runs, so the openers below are the ones
# `from pptx import Presentation` and the like get.
_opened = {}


def _fingerprint(document):
    digest = hashlib.sha1()
    if hasattr(document, "slides"):  # python-pptx: every slide's XML, in order
        for slide in document.slides:
            digest.update(etree.tostring(slide._element))
    elif hasattr(document, "worksheets"):  # openpyxl: sheet names and cell values
        for sheet in document.worksheets:
            digest.update(repr((sheet.title, [[c.value for c in row] for row in sheet.iter_rows()])).encode())
    elif hasattr(document, "element"):  # python-docx: the document body
        digest.update(etree.tostring(document.element))
    return digest.hexdigest()


def _tracking(opener):
    def open_document(source=None, *args, **kwargs):
        document = opener(source, *args, **kwargs)
        if isinstance(source, str):
            _opened[id(document)] = (source, _fingerprint(document))
        return document

    open_document.__doc__ = opener.__doc__
    return open_document


if not getattr(pptx, "_katechat_tracked", False):
    pptx.Presentation = _tracking(pptx.Presentation)
    docx.Document = _tracking(docx.Document)
    openpyxl.load_workbook = _tracking(openpyxl.load_workbook)
    pptx._katechat_tracked = True


def _paragraphs(container):
    """Every paragraph of a python-docx document or a python-pptx presentation, tables included."""
    if hasattr(container, "slides"):  # a presentation
        for slide in container.slides:
            yield from _slide_paragraphs(slide)
        return
    parts = [container]
    for section in getattr(container, "sections", []):
        parts += [section.header, section.footer]
    for part in parts:
        yield from part.paragraphs
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs


def _slide_paragraphs(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield from shape.text_frame.paragraphs
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs
    if slide.has_notes_slide:
        yield from slide.notes_slide.notes_text_frame.paragraphs


def _a(tag):
    return "{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag


def _paragraph_level(paragraph):
    properties = paragraph.find(_a("pPr"))
    return int(properties.get("lvl", "0")) if properties is not None else 0


def _shrink_to_fit(shape, body, lines):
    """Stores a shrink-to-fit scale (PowerPoint draws with it) when the text will not fit the shape."""
    if not getattr(shape, "width", None) or not getattr(shape, "height", None):
        return
    sizes = [int(r.get("sz")) / 100 for r in body.iter(_a("rPr")) if r.get("sz")]
    size = sizes[0] if sizes else (24 if getattr(shape, "is_placeholder", False) else 18)
    width, height = shape.width / 12700, shape.height / 12700  # EMU → points
    chars_per_line = max(1.0, width / (size * 0.5))
    lines_that_fit = max(1.0, height / (size * 1.2))
    needed = sum(max(1, -(-len(line) // int(chars_per_line))) for line in lines)
    if needed <= lines_that_fit:
        return
    properties = body.find(_a("bodyPr"))
    if properties is None:
        return
    for fit in ("noAutofit", "spAutoFit", "normAutofit"):
        for element in properties.findall(_a(fit)):
            properties.remove(element)
    scale = max(0.5, (lines_that_fit / needed) ** 0.5)
    fit = properties.makeelement(_a("normAutofit"), {"fontScale": str(int(scale * 100000))})
    warp = properties.find(_a("prstTxWarp"))
    if warp is not None:
        warp.addnext(fit)
    else:
        properties.insert(0, fit)


def set_text(target, text):
    """Replaces ALL the text of a shape (or text frame) or a Word paragraph with `text`.

    The old text is removed, not appended to. `text` is a string, or a list with one item per
    paragraph (bullets of a body placeholder); an item may be (text, level) for a sub-bullet.
    New paragraphs take the look (bullet, font, size, colour) of the shape's first paragraph, or of
    its first paragraph at that level. On a slide, text too long for its shape is shrunk to fit.
    """
    if hasattr(target, "runs") and hasattr(target, "style"):  # a python-docx paragraph
        runs = target.runs
        for run in runs[1:]:
            run._element.getparent().remove(run._element)
        items = text if isinstance(text, (list, tuple)) else [text]
        value = "\n".join(str(item[0] if isinstance(item, tuple) else item) for item in items)
        if runs:
            runs[0].text = value
        else:
            target.add_run(value)
        return target

    if hasattr(target, "has_text_frame") and not target.has_text_frame:
        slide = getattr(getattr(target, "part", None), "slide", None)
        shapes = _describe(slide) if slide is not None else "?"
        raise ValueError(
            f"set_text: {target.name!r} is a {_role(target)} and holds no text. "
            f"The shapes of its slide are (index, role, name, text): {shapes}; "
            "pick a text shape with shape_by_role(slide, 'title' | 'subtitle' | 'body')"
        )
    frame = target.text_frame if hasattr(target, "text_frame") else target
    body = frame._txBody
    items = list(text) if isinstance(text, (list, tuple)) else str(text).split("\n")
    old_paragraphs = body.findall(_a("p"))

    new_paragraphs, lines = [], []
    for item in items or [""]:
        line, level = (str(item[0]), int(item[1])) if isinstance(item, tuple) else (str(item), None)
        lines.append(line)
        template = None
        if level is not None:
            template = next((p for p in old_paragraphs if _paragraph_level(p) == level), None)
        if template is None and old_paragraphs:
            template = old_paragraphs[0]
        paragraph = copy.deepcopy(template) if template is not None else body.makeelement(_a("p"), {})
        runs = paragraph.findall(_a("r"))
        style = runs[0].find(_a("rPr")) if runs else paragraph.find(_a("endParaRPr"))
        for child in list(paragraph):
            if child.tag not in (_a("pPr"), _a("endParaRPr")):
                paragraph.remove(child)
        if level is not None:
            properties = paragraph.find(_a("pPr"))
            if properties is None:
                properties = paragraph.makeelement(_a("pPr"), {})
                paragraph.insert(0, properties)
            properties.set("lvl", str(level))
        run = paragraph.makeelement(_a("r"), {})
        if style is not None:
            run_style = copy.deepcopy(style)
            run_style.tag = _a("rPr")
            run.append(run_style)
        text_element = paragraph.makeelement(_a("t"), {})
        text_element.text = line
        run.append(text_element)
        end = paragraph.find(_a("endParaRPr"))
        if end is not None:
            end.addprevious(run)
        else:
            paragraph.append(run)
        new_paragraphs.append(paragraph)

    for paragraph in old_paragraphs:
        body.remove(paragraph)
    for paragraph in new_paragraphs:
        body.append(paragraph)
    if hasattr(target, "text_frame"):
        _shrink_to_fit(target, body, lines)
    return target


def _role(shape):
    """title, subtitle, body, picture, table, chart or other"""
    kind = str(shape.placeholder_format.type).lower() if shape.is_placeholder else ""
    if "picture" in kind or (shape.shape_type is not None and "PICTURE" in str(shape.shape_type)):
        return "picture"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if not shape.has_text_frame:
        return "other"
    if "subtitle" in kind:
        return "subtitle"
    if "title" in kind:
        return "title"
    return "body" if shape.is_placeholder or shape.text_frame.text.strip() else "other"


def _describe(slide):
    return [
        (position, _role(shape), shape.name, shape.text_frame.text if shape.has_text_frame else "")
        for position, shape in enumerate(slide.shapes)
    ]


def slide_shapes(prs, index):
    """What slide `index` holds, to pick the shapes to change: [(shape index, role, name, text)].

    role is "title", "subtitle", "body", "picture", "table", "chart" or "other".
    """
    return _describe(prs.slides[index])


def shape_by_role(slide, role, nth=0):
    """The slide's `nth` shape (from 0, top to bottom in the slide's order) with this role: "title",
    "subtitle" or "body" for text, "picture", "table" or "chart". Raises, listing what the slide
    has, when there is none."""
    matches = [shape for shape in slide.shapes if _role(shape) == role]
    if nth < len(matches):
        return matches[nth]
    raise ValueError(f"No {role} shape #{nth} on this slide. It has: {_describe(slide)}")


def delete_shape(shape):
    """Removes a shape (text box, picture, ...) from its slide."""
    shape._element.getparent().remove(shape._element)


def delete_paragraph(paragraph):
    """Removes a paragraph from a Word document (or a slide's text)."""
    element = paragraph._element if hasattr(paragraph, "_element") else paragraph._p
    element.getparent().remove(element)


def replace_text(container, old, new):
    """Replaces text everywhere in a document or presentation, keeping the formatting.

    Text split across runs (e.g. partly bold) is found too; the replaced paragraph then takes the
    formatting of its first run. Returns how many paragraphs changed.
    """
    changed = 0
    for paragraph in _paragraphs(container):
        runs = list(paragraph.runs)
        if not runs:
            continue
        text = "".join(run.text for run in runs)
        if old not in text:
            continue
        # prefer a replacement inside single runs, which keeps mixed formatting
        if all(old not in run.text for run in runs) or text.count(old) != sum(run.text.count(old) for run in runs):
            runs[0].text = text.replace(old, new)
            for run in runs[1:]:
                run.text = ""
        else:
            for run in runs:
                run.text = run.text.replace(old, new)
        changed += 1
    if not changed:
        print(f"replace_text: {old!r} was not found, nothing replaced")
    return changed


def set_font(container, name, size_pt=None):
    """Sets the font (and optionally the size in points) of all text in a document or presentation."""
    from docx.shared import Pt as DocxPt
    from pptx.util import Pt as PptxPt

    is_pptx = hasattr(container, "slides")
    for paragraph in _paragraphs(container):
        for run in paragraph.runs:
            run.font.name = name
            if size_pt:
                run.font.size = PptxPt(size_pt) if is_pptx else DocxPt(size_pt)
    if not is_pptx:
        for style in container.styles:
            if getattr(style, "font", None) is not None:
                style.font.name = name


def set_text_color(container, hex_color, headings_only=False):
    """Colours all text ("1F2933"), or only titles/headings with headings_only=True."""
    from docx.shared import RGBColor as DocxRGB

    is_pptx = hasattr(container, "slides")
    if is_pptx:
        for slide in container.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                is_title = shape.is_placeholder and "title" in str(shape.placeholder_format.type).lower()
                if headings_only and not is_title:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = PptxRGB.from_string(hex_color)
        return
    for paragraph in _paragraphs(container):
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        if headings_only and not (style.startswith("heading") or style == "title"):
            continue
        for run in paragraph.runs:
            run.font.color.rgb = DocxRGB.from_string(hex_color)


def set_background(prs, hex_color):
    """Gives every slide of a presentation a solid background colour ("FFF8E1")."""
    for slide in prs.slides:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = PptxRGB.from_string(hex_color)


def slide_texts(prs):
    """The text of each slide, for finding the one to change: [(index, "title\\nbody...")]."""
    return [
        (index, "\n".join(p.text for p in _slide_paragraphs(slide) if p.text.strip()))
        for index, slide in enumerate(prs.slides)
    ]


def duplicate_slide(prs, index):
    """Appends a copy of slide `index` (shapes, pictures, charts, background) and returns it.

    Move it into place with move_slide. A copied chart shares its data with the original.
    """
    source = prs.slides[index]
    duplicate = prs.slides.add_slide(source.slide_layout)
    for shape in list(duplicate.shapes):
        shape._element.getparent().remove(shape._element)

    # pictures, charts and links point at parts through relationship ids: give the copy the same
    # relationships and rewrite the ids in the copied shapes
    ids = {}
    for rel_id, rel in source.part.rels.items():
        if rel.reltype.endswith("/notesSlide") or rel.reltype.endswith("/slideLayout"):
            continue
        if rel.is_external:
            ids[rel_id] = duplicate.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            ids[rel_id] = duplicate.part.rels.get_or_add(rel.reltype, rel.target_part)

    def remap(element):
        for node in element.iter():
            for name, value in list(node.attrib.items()):
                if name.rsplit("}", 1)[-1] in ("embed", "id", "link", "pict") and value in ids:
                    node.set(name, ids[value])
        return element

    for shape in source.shapes:
        duplicate.shapes._spTree.insert_element_before(remap(copy.deepcopy(shape._element)), "p:extLst")
    background = source._element.cSld.bg
    if background is not None:
        duplicate._element.cSld.insert(0, remap(copy.deepcopy(background)))
    return duplicate


def delete_slide(prs, index):
    """Removes slide `index`."""
    slide_ids = prs.slides._sldIdLst
    slide_id = slide_ids[index]
    prs.part.drop_rel(slide_id.rId)
    slide_ids.remove(slide_id)


def move_slide(prs, old_index, new_index):
    """Moves a slide to another position (indexes from 0; -1 is the end)."""
    slide_ids = prs.slides._sldIdLst
    slide_id = slide_ids[old_index]
    slide_ids.remove(slide_id)
    if new_index < 0:
        new_index = len(slide_ids) + 1 + new_index
    slide_ids.insert(new_index, slide_id)


def save(document, file_name):
    """Saves a python-docx document, python-pptx presentation or openpyxl workbook to /output/<file_name>."""
    opened = _opened.get(id(document))
    if opened and _fingerprint(document) == opened[1]:
        raise RuntimeError(
            f"Nothing was changed: {os.path.basename(file_name)} would be the same as {opened[0]}. "
            "Make the changes the user asked for: replace_text returns how many paragraphs it changed "
            "(0 means the text is not in the document; slide_texts(prs) shows the text a deck has)."
        )
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    path = os.path.join(OUTPUT_DIR, os.path.basename(file_name))
    document.save(path)
    return path
